import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from party_matching.reporting import _metrics, _rules_demo_rows, _write_workbook


class RulesDemoReportTests(unittest.TestCase):
    def test_rules_projection_preserves_ml_and_counts_no_match_in_recall(self):
        base = {
            "Dataset Split": "TEST_UNSEEN", "Held-Out": True,
            "Scorable": True, "Evaluation Status": "LABELED_HELD_OUT",
            "Category": "Corporate", "Has OBO": False, "Has VIA": False,
            "Expected Canonical Name": "Acme Corp", "Expected Global Parent Name": "Acme Corp",
            "Expected Verified ID": "acme", "Decision": "NO_MATCH", "Correct": False,
            "Reason": "INSUFFICIENT_SUPPORT", "Confidence": 0.2,
            "Predicted Verified ID": None, "Predicted Verified Name": None,
            "Expected Target Retrieved": False, "Expected Target Rank": None,
            "Error Bucket": "RETRIEVAL_MISS", "Case Types": "GENERAL",
            "Connector Resolution": None, "Rules-only Verified ID": None,
            "Rules-only Decision Tier": None, "Prior Rules Verified ID": None,
            "Prior Rules Decision Tier": None,
        }
        first = dict(base, **{"ADM Party ID": "one"})
        second = dict(base, **{"ADM Party ID": "two"})
        rows = [
            {"source": {"Raw": "Acme"}, "prediction": first},
            {"source": {"Raw": "Unresolved"}, "prediction": second},
        ]
        rules = {
            "one": {"decision": "MATCH", "verified_party_id": "acme",
                    "verified_party_name": "Acme Corp", "confidence": 0.95,
                    "reason": "MATCH", "retrieved_root_ids": ["acme"]},
            "two": {"decision": "NO_MATCH", "confidence": 0.35,
                    "reason": "INSUFFICIENT_SUPPORT", "retrieved_root_ids": ["acme"]},
        }

        projected = _rules_demo_rows(rows, rules, {"raw_name_column": "Raw"})
        result = _metrics(projected, {})["overall"]

        self.assertEqual(rows[0]["prediction"]["Decision"], "NO_MATCH")
        self.assertEqual(projected[0]["prediction"]["Predicted Verified Name"], "Acme Corp")
        self.assertEqual(projected[1]["prediction"]["Decision"], "NO_MATCH")
        self.assertEqual(result["scorable"], 2)
        self.assertEqual(result["correct_matches"], 1)
        self.assertEqual(result["recall"], 0.5)

    def test_all_labeled_recall_excludes_error_and_unknown_but_keeps_no_match(self):
        def row(split, category, decision, correct):
            held_out = split in {"TEST_KNOWN", "TEST_UNSEEN"}
            scorable = split != "UNKNOWN"
            prediction = {
                "Dataset Split": split, "Held-Out": held_out,
                "Scorable": scorable,
                "Evaluation Status": "UNKNOWN" if not scorable else "LABELED_HELD_OUT" if held_out else "LABELED_DEVELOPMENT",
                "Category": category, "Has OBO": False, "Has VIA": False,
                "Case Types": "GENERAL", "Expected Canonical Name": "Acme Corp",
                "Expected Verified ID": "acme", "Decision": decision,
                "Correct": correct, "Confidence": 0.9 if decision == "MATCH" else 0.2,
                "Reason": "MATCH" if decision == "MATCH" else "INSUFFICIENT_SUPPORT",
                "Predicted Verified ID": "acme" if correct else None,
                "Predicted Verified Name": "Acme Corp" if correct else None,
                "Expected Target Retrieved": True, "Expected Target Rank": 1,
                "Error Bucket": "NONE" if correct else "REJECTED_AFTER_RETRIEVAL",
                "Connector Resolution": None,
                "Rules-only Verified ID": "acme" if correct else None,
                "Rules-only Decision Tier": None,
                "Prior Rules Verified ID": None,
                "Prior Rules Decision Tier": None,
            }
            return {"source": {"Raw": "Example"}, "prediction": prediction}

        rows = [
            row("TRAIN", "Corporate", "MATCH", True),
            row("CALIBRATION", "Corporate", "NO_MATCH", False),
            row("TEST_UNSEEN", "Corporate", "MATCH", True),
            row("TEST_KNOWN", " ERROR ", "NO_MATCH", False),
            row("UNKNOWN", "UNKNOWN", "NO_MATCH", None),
        ]
        metrics = _metrics(rows, {})
        all_labeled = metrics["all_labeled_excluding_error"]
        self.assertEqual(all_labeled["scorable"], 3)
        self.assertEqual(all_labeled["correct_matches"], 2)
        self.assertAlmostEqual(all_labeled["recall"], 2 / 3)
        self.assertEqual(metrics["overall"]["scorable"], 2)
        self.assertEqual(metrics["overall"]["recall"], 0.5)
        self.assertEqual(metrics["all_labeled"]["scorable"], 4)

        with TemporaryDirectory() as directory:
            output = Path(directory) / "predictions_rules_only.xlsx"
            _write_workbook(
                output, rows,
                {"raw_name_column": "Raw", "rows": len(rows), "cleaned_unverified_rows": len(rows)},
                metrics, {}, view="rules",
            )
            workbook = load_workbook(output, read_only=True, data_only=True)
            stats = workbook["Stats"]
            stats_values = {
                first.value: second.value
                for first, second in stats.iter_rows(min_col=1, max_col=2)
                if first.value
            }
            self.assertEqual(stats_values["All labeled rows (excl. ERROR)"], 3)
            self.assertEqual(stats_values["All labeled correct (excl. ERROR)"], 2)
            self.assertAlmostEqual(stats_values["All labeled recall (excl. ERROR)"], 2 / 3)
            self.assertEqual(stats["D7"].value, "All labeled excl. ERROR")
            self.assertAlmostEqual(stats["H7"].value, 2 / 3)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
