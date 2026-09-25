import json
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from suggestions.family_demo import build_family_demo, _save_stage
from party_matching.domain import load_config


def test_full_population_independent_initial_and_audited_final():
    with TemporaryDirectory() as temporary:
        root = Path(temporary)
        prepared = root / "prepared"
        prepared.mkdir()
        catalog = [
            {"partyId": "a", "partyName": "Abbott Laboratories"},
            {"partyId": "n", "partyName": "Abbott Nutrition"},
            {"partyId": "d", "partyName": "DocuSign, Inc."},
            {"partyId": "o", "partyName": "Other Party"},
        ]
        (prepared / "verified_parties.json").write_text(json.dumps(catalog), encoding="utf-8")
        (prepared / "jobs.json").write_text(json.dumps([{
            "accountId": "test", "confidenceCutoff": 0.0, "verifiedParties": catalog,
        }]), encoding="utf-8")
        names = ["Abbott Laboratories", "Abbott Nutrition", "DocuSign Brazil",
                 "Unrelated Name", "DocuSign, Inc. OBO Other Party",
                 "Other Party OBO DocuSign, Inc."]
        (prepared / "adm_records.jsonl").write_text(
            "\n".join(json.dumps({"account_id": "test", "adm_party_id": str(i),
                                   "raw_name": name, "source_row": i,
                                   "eligible": True, "is_verified": False})
                      for i, name in enumerate(names, 1)) + "\n", encoding="utf-8")
        (prepared / "labels.jsonl").write_text(
            "\n".join(json.dumps({"adm_party_id": str(i), "expected_party_id": party_id,
                                   "expected_canonical_name": expected,
                                   "split": "TEST_UNSEEN", "scorable": True})
                      for i, party_id, expected in [
                          (1, "a", "Abbott Laboratories"),
                          (2, "n", "Abbott Nutrition"),
                          (4, "a", "Abbott Laboratories"),
                      ]) + "\n", encoding="utf-8")
        config = load_config("config.toml")
        config["paths"]["state_dir"] = str(root / "state")
        config["paths"]["artifacts_dir"] = str(root / "artifacts")
        config["execution"]["mode"] = "serial"
        config["retrieval"]["embedding_enabled"] = False
        initial, final, common = build_family_demo(
            prepared, config, [{"name": "Abbott Laboratories"}, {"name": "DocuSign, Inc."}])

        assert common["eligible_unverified_rows"] == len(names)
        assert all(party["eligible_unverified"] == len(names) for party in initial["parties"])
        assert any(row["verified_party"] == "Abbott Laboratories"
                   and row["unverified_party"] == "Abbott Nutrition"
                   for row in initial["review"] + initial["dropped"])
        assert not any(row["verified_party"] == "Abbott Laboratories"
                       and row["unverified_party"] == "Abbott Nutrition"
                       for row in initial["strong"])
        assert any(row["unverified_party"] == "Unrelated Name"
                   and row["verified_party"] == "Abbott Laboratories"
                   for row in initial["misses"])
        assert all(row["reason"] != "ASSIGNED_TO_OTHER_PARTY" for row in initial["dropped"])
        assert any(row["verified_party"] == "Abbott Laboratories"
                   and row["unverified_party"] == "Abbott Nutrition"
                   and row["reason"] == "ASSIGNED_TO_OTHER_PARTY"
                   and row["assigned_to"] == "Abbott Nutrition"
                   for row in final["dropped"])
        assert not any(row["verified_party"] == "DocuSign, Inc."
                       and row["unverified_party"] == "Other Party OBO DocuSign, Inc."
                       for row in initial["strong"] + initial["review"] + initial["dropped"])
        assert all(row["verified_party_id"] in {"a", "d"} for row in final["strong"])
        assert not (root / "state").exists()

        folder = root / "initial"
        _save_stage(folder, initial)
        book = load_workbook(folder / "suggestions_current.xlsx", read_only=True)
        try:
            assert book.sheetnames == ["Verified parties", "Strong suggestions",
                                      "Review candidates", "Dropped candidates",
                                      "Known-label gaps", "Stats"]
            headers = [cell.value for cell in next(book["Review candidates"].rows)]
            assert "Matched verified party" in headers
        finally:
            book.close()
        saved = json.loads((folder / "metrics.json").read_text(encoding="utf-8"))
        assert saved["eligible_unverified_rows"] == len(names)
